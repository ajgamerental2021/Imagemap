"""Render uploaded templates and simple web-designed templates."""

from __future__ import annotations

import io
import re
from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from keyword_engine import PLACEHOLDER, substitute, value_map

ROOT = Path(__file__).parent
ALLOWED_BLOCKS = {"heading", "paragraph", "field", "divider"}


def web_spec(raw: dict) -> dict:
    if not isinstance(raw, dict):
        raise ValueError("รูปแบบ template หน้าเว็บไม่ถูกต้อง")
    title = str(raw.get("title", "เอกสารใหม่")).strip()[:200]
    accent = str(raw.get("accent", "#2d7655"))
    if not re.fullmatch(r"#[0-9a-fA-F]{6}", accent):
        raise ValueError("สีหลักต้องเป็นรหัส #RRGGBB")
    try:
        font_size = int(raw.get("font_size", 12))
    except (TypeError, ValueError) as exc:
        raise ValueError("ขนาดตัวอักษรไม่ถูกต้อง") from exc
    if not 9 <= font_size <= 22:
        raise ValueError("ขนาดตัวอักษรต้องอยู่ระหว่าง 9 ถึง 22")
    blocks = raw.get("blocks", [])
    if not isinstance(blocks, list) or len(blocks) > 500:
        raise ValueError("template ต้องมีบล็อกไม่เกิน 500 รายการ")
    cleaned = []
    for block in blocks:
        if not isinstance(block, dict) or block.get("type") not in ALLOWED_BLOCKS:
            raise ValueError("ชนิดบล็อกใน template ไม่ถูกต้อง")
        cleaned.append({"type": block["type"], "text": str(block.get("text", ""))[:10000],
                        "label": str(block.get("label", ""))[:200]})
    return {"title": title or "เอกสารใหม่", "accent": accent, "font_size": font_size, "blocks": cleaned}


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def render_web_docx(spec: dict, values: dict) -> bytes:
    doc = Document()
    doc.sections[0].top_margin = Inches(0.7)
    doc.sections[0].bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Noto Sans Thai"
    doc.styles["Normal"].font.size = Pt(spec["font_size"])
    title = doc.add_heading(substitute(spec["title"], values), 0)
    for run in title.runs:
        run.font.color.rgb = rgb(spec["accent"])
    for block in spec["blocks"]:
        block_type = block["type"]
        if block_type == "divider":
            doc.add_paragraph("────────────────────────────────")
        elif block_type == "heading":
            paragraph = doc.add_heading(substitute(block["text"], values), 1)
            for run in paragraph.runs:
                run.font.color.rgb = rgb(spec["accent"])
        elif block_type == "paragraph":
            doc.add_paragraph(substitute(block["text"], values))
        elif block_type == "field":
            paragraph = doc.add_paragraph()
            paragraph.add_run(substitute(block["label"], values) + ": ").bold = True
            paragraph.add_run(substitute(block["text"], values))
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def render_web_xlsx(spec: dict, values: dict) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Document"
    sheet.column_dimensions["A"].width = 28
    for letter in "BCD":
        sheet.column_dimensions[letter].width = 22
    accent = spec["accent"].lstrip("#")
    row = 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    title = sheet.cell(row, 1, substitute(spec["title"], values))
    title.font = Font(name="Noto Sans Thai", size=spec["font_size"] + 6, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=accent)
    title.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 38
    for block in spec["blocks"]:
        row += 1
        block_type = block["type"]
        if block_type == "divider":
            sheet.row_dimensions[row].height = 9
            for column in range(1, 5):
                sheet.cell(row, column).fill = PatternFill("solid", fgColor=accent)
        elif block_type == "field":
            lines = substitute(block["text"], values).split("\n")
            for line_number, line in enumerate(lines):
                if len(line) > 32767:
                    raise ValueError("ค่าหนึ่งรายการยาวเกินขนาดเซลล์ Excel 32,767 ตัวอักษร")
                target_row = row + line_number
                if line_number == 0:
                    sheet.cell(target_row, 1, substitute(block["label"], values)).font = Font(name="Noto Sans Thai", bold=True, size=spec["font_size"])
                sheet.merge_cells(start_row=target_row, start_column=2, end_row=target_row, end_column=4)
                value = sheet.cell(target_row, 2, line)
                value.alignment = Alignment(wrap_text=True, vertical="top")
                value.font = Font(name="Noto Sans Thai", size=spec["font_size"])
                sheet.row_dimensions[target_row].height = 36
            row += len(lines) - 1
        else:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            rendered = substitute(block["text"], values)
            if len(rendered) > 32767:
                raise ValueError("ข้อความยาวเกินขนาดเซลล์ Excel 32,767 ตัวอักษร")
            value = sheet.cell(row, 1, rendered)
            value.alignment = Alignment(wrap_text=True, vertical="top")
            value.font = Font(name="Noto Sans Thai", size=spec["font_size"] + (2 if block_type == "heading" else 0),
                              bold=block_type == "heading", color=accent if block_type == "heading" else "24323A")
            sheet.row_dimensions[row].height = 38 if block_type == "heading" else 52
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def register_pdf_fonts():
    if "NotoThai" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("NotoThai", str(ROOT / "assets" / "NotoSansThai-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("NotoThai-Bold", str(ROOT / "assets" / "NotoSansThai-Bold.ttf")))
        pdfmetrics.registerFontFamily("NotoThai", normal="NotoThai", bold="NotoThai-Bold")


def pdf_paragraph(content: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(content).replace("\n", "<br/>"), style)


def render_web_pdf(spec: dict, values: dict) -> bytes:
    register_pdf_fonts()
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48)
    accent = colors.HexColor(spec["accent"])
    base = ParagraphStyle("base", fontName="NotoThai", fontSize=spec["font_size"],
                          leading=spec["font_size"] * 1.65, alignment=TA_LEFT, spaceAfter=11)
    title_style = ParagraphStyle("title", parent=base, fontName="NotoThai-Bold", fontSize=spec["font_size"] + 8,
                                 leading=(spec["font_size"] + 8) * 1.5, textColor=accent, spaceAfter=18)
    heading_style = ParagraphStyle("heading", parent=base, fontName="NotoThai-Bold", fontSize=spec["font_size"] + 3,
                                   leading=(spec["font_size"] + 3) * 1.5, textColor=accent)
    story = [pdf_paragraph(substitute(spec["title"], values), title_style)]
    for block in spec["blocks"]:
        block_type = block["type"]
        if block_type == "divider":
            story.extend([Spacer(1, 5), HRFlowable(width="100%", color=accent, thickness=1), Spacer(1, 12)])
        elif block_type == "heading":
            story.append(pdf_paragraph(substitute(block["text"], values), heading_style))
        elif block_type == "paragraph":
            story.append(pdf_paragraph(substitute(block["text"], values), base))
        elif block_type == "field":
            story.append(pdf_paragraph(substitute(block["label"], values), heading_style))
            story.append(pdf_paragraph(substitute(block["text"], values), base))
    document.build(story)
    return output.getvalue()


def render_web_template(raw_spec: dict, values: dict, output_format: str) -> tuple[bytes, str]:
    spec = web_spec(raw_spec)
    value_map(values)
    if output_format == "docx":
        return render_web_docx(spec, values), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if output_format == "xlsx":
        return render_web_xlsx(spec, values), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if output_format == "pdf":
        return render_web_pdf(spec, values), "application/pdf"
    raise ValueError("รองรับผลลัพธ์ DOCX, XLSX หรือ PDF")


def replace_in_paragraph(paragraph, values: dict) -> int:
    runs = paragraph.runs
    if not runs:
        return 0
    lengths = [len(run.text) for run in runs]
    full = "".join(run.text for run in runs)
    matches = list(PLACEHOLDER.finditer(full))
    replaced = 0
    for match in reversed(matches):
        replacement = substitute(match.group(0), values)
        if replacement == match.group(0):
            continue
        start, end = match.span()
        first = last = None
        offset = 0
        for index, length in enumerate(lengths):
            if first is None and start < offset + length:
                first = (index, start - offset)
            if end <= offset + length:
                last = (index, end - offset)
                break
            offset += length
        if first is None or last is None:
            continue
        first_index, first_offset = first
        last_index, last_offset = last
        if first_index == last_index:
            text = runs[first_index].text
            runs[first_index].text = text[:first_offset] + replacement + text[last_offset:]
        else:
            runs[first_index].text = runs[first_index].text[:first_offset] + replacement
            for index in range(first_index + 1, last_index):
                runs[index].text = ""
            runs[last_index].text = runs[last_index].text[last_offset:]
        replaced += 1
    return replaced


def fill_docx(data: bytes, values: dict) -> bytes:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("อ่านไฟล์ template Word ไม่สำเร็จ") from exc
    count = 0

    def walk(container):
        nonlocal count
        for paragraph in container.paragraphs:
            count += replace_in_paragraph(paragraph, values)
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    walk(cell)

    walk(doc)
    for section in doc.sections:
        walk(section.header)
        walk(section.footer)
    if not count:
        raise ValueError("ไม่พบ placeholder {{keyword}} ที่มีข้อมูลใน template Word")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def fill_xlsx(data: bytes, values: dict, keep_vba: bool = False) -> bytes:
    try:
        book = load_workbook(io.BytesIO(data), keep_vba=keep_vba)
    except Exception as exc:
        raise ValueError("อ่านไฟล์ template Excel ไม่สำเร็จ") from exc
    count = 0
    for sheet in book:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    rendered = substitute(cell.value, values)
                    if rendered != cell.value:
                        if len(rendered) > 32767:
                            raise ValueError(f"{sheet.title}!{cell.coordinate} ยาวเกินขนาดเซลล์ Excel 32,767 ตัวอักษร")
                        cell.value = rendered
                        count += 1
    if not count:
        raise ValueError("ไม่พบ placeholder {{keyword}} ที่มีข้อมูลใน template Excel")
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def fill_pdf_form(data: bytes, values: dict) -> bytes:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("อ่านไฟล์ template PDF ไม่สำเร็จ") from exc
    fields = reader.get_fields() or {}
    if not fields:
        raise ValueError("PDF template ต้องเป็นแบบฟอร์มที่มีช่องชื่อเดียวกับ keyword")
    replacements = {}
    for name in fields:
        rendered = substitute("{{" + name + "}}", values)
        if rendered != "{{" + name + "}}":
            replacements[name] = rendered
    if not replacements:
        raise ValueError("ไม่พบช่อง PDF ที่ชื่อเดียวกับ keyword ที่มีข้อมูล")
    writer = PdfWriter()
    writer.append(reader)
    writer.update_page_form_field_values(None, replacements, auto_regenerate=False)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def render_file_template(filename: str, data: bytes, values: dict) -> tuple[bytes, str, str]:
    suffix = Path(filename).suffix.lower()
    value_map(values)
    if suffix == ".docx":
        return fill_docx(data, values), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    if suffix in {".xlsx", ".xlsm"}:
        return fill_xlsx(data, values, keep_vba=suffix == ".xlsm"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", suffix[1:]
    if suffix == ".pdf":
        return fill_pdf_form(data, values), "application/pdf", "pdf"
    raise ValueError("template ที่อัปโหลดรองรับ DOCX, XLSX, XLSM หรือ PDF แบบฟอร์ม")
