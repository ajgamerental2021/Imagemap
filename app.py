"""Local, stateless document and image comparison API."""

from __future__ import annotations

import difflib
import csv
import io
import json
import os
import re
import subprocess
import sys
import tempfile
from math import sqrt
from itertools import accumulate
from itertools import zip_longest
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, Response
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageOps, UnidentifiedImageError
from pypdf import PdfReader
from pypdfium2 import PdfDocument
from pyxlsb import open_workbook as open_xlsb
from pydantic import BaseModel
from starlette.concurrency import run_in_threadpool
from xlrd import open_workbook as open_xls

from document_data import SPREADSHEET_TYPES, cell_differences, iter_cell_changes, read_spreadsheet
from generation import render_file_template, render_web_template
from keyword_engine import extract_hits, parse_keywords

ROOT = Path(__file__).parent
MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 200_000
MAX_IMAGE_PIXELS = 18_000_000
MAX_PDF_PAGES = 100
MAX_PDF_OCR_PAGES = 30
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp", ".avif", ".bmp", ".gif", ".tif", ".tiff"}
TEXT_TYPES = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".yaml", ".yml", ".log", ".py", ".js", ".ts", ".css", ".sql"}
DOCUMENT_TYPES = {".pdf", ".docx", ".xlsx", ".xlsm", ".xls", ".xlsb"}
SUPPORTED_TYPES = IMAGE_TYPES | TEXT_TYPES | DOCUMENT_TYPES

app = FastAPI(title="Pairwise Compare", docs_url=None, redoc_url=None)


@app.get("/")
def index():
    return FileResponse(ROOT / "static" / "index.html")


@app.get("/static/{filename}")
def static(filename: str):
    if filename not in {"app.js", "generator.js", "styles.css"}:
        raise HTTPException(404)
    return FileResponse(ROOT / "static" / filename)


@app.get("/api/health")
def health():
    return {"ok": True}


@app.get("/api/capabilities")
def capabilities():
    enabled = os.environ.get("LOCAL_FOLDER_ACCESS") == "1"
    return {"local_folder_access": enabled, "native_folder_picker": enabled and sys.platform == "darwin"}


def kind(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in IMAGE_TYPES:
        return "image"
    if suffix in TEXT_TYPES | DOCUMENT_TYPES:
        return "document"
    raise ValueError(f"ไม่รองรับไฟล์ชนิด {suffix or '(ไม่มีนามสกุล)'}")


def spreadsheet_value(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


async def read_upload(upload: UploadFile) -> bytes:
    data = await upload.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ValueError("ไฟล์ใหญ่เกิน 20 MB")
    if not data:
        raise ValueError("ไฟล์ว่าง")
    return data


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in TEXT_TYPES:
        try:
            decoded = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("ไฟล์ข้อความต้องเข้ารหัส UTF-8") from exc
        if "\x00" in decoded:
            raise ValueError("ไฟล์นี้ดูเป็นข้อมูลไบนารี ไม่ใช่ข้อความ")
        result = decoded
    elif suffix == ".pdf":
        result, _ = extract_pdf_text(data)
    elif suffix == ".docx":
        try:
            doc = Document(io.BytesIO(data))
            sections = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                sections.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
            result = "\n".join(sections)
        except Exception as exc:
            raise ValueError("อ่าน DOCX ไม่สำเร็จ") from exc
    elif suffix in {".xlsx", ".xlsm"}:
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sections = []
            for sheet in book:
                sections.append(f"[ชีต: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    sections.append("\t".join(spreadsheet_value(cell) for cell in row))
            book.close()
            result = "\n".join(sections)
        except Exception as exc:
            raise ValueError("อ่านไฟล์ Excel ไม่สำเร็จ") from exc
    elif suffix == ".xls":
        try:
            book = open_xls(file_contents=data)
            sections = []
            for sheet in book.sheets():
                sections.append(f"[ชีต: {sheet.name}]")
                for row_number in range(sheet.nrows):
                    sections.append("\t".join(spreadsheet_value(cell) for cell in sheet.row_values(row_number)))
            result = "\n".join(sections)
        except Exception as exc:
            raise ValueError("อ่าน XLS ไม่สำเร็จ") from exc
    elif suffix == ".xlsb":
        try:
            sections = []
            with open_xlsb(io.BytesIO(data)) as book:
                for name in book.sheets:
                    sections.append(f"[ชีต: {name}]")
                    with book.get_sheet(name) as sheet:
                        for row in sheet.rows():
                            sections.append("\t".join(spreadsheet_value(cell.v) for cell in row))
            result = "\n".join(sections)
        except Exception as exc:
            raise ValueError("อ่าน XLSB ไม่สำเร็จ") from exc
    else:
        raise ValueError("ไม่รองรับชนิดเอกสารนี้")
    if len(result) > MAX_TEXT:
        raise ValueError("ข้อความที่ดึงได้ยาวเกิน 200,000 ตัวอักษร")
    return result


def load_image(data: bytes) -> Image.Image:
    try:
        image = Image.open(io.BytesIO(data))
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise ValueError("รูปมีพิกเซลมากเกิน 18 ล้านพิกเซล")
        return ImageOps.exif_transpose(image).convert("RGBA")
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise ValueError("อ่านรูปไม่สำเร็จ") from exc


def ocr_image(image: Image.Image) -> str | None:
    """Return None if OCR is unavailable, and text (possibly empty) otherwise."""
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / "image.png"
        image.convert("RGB").save(source)
        try:
            environment = os.environ.copy()
            tessdata = ROOT / ".local" / "tessdata"
            if (tessdata / "tha.traineddata").is_file():
                environment["TESSDATA_PREFIX"] = str(tessdata)
            langs = subprocess.run(
                ["tesseract", "--list-langs"], capture_output=True, text=True, timeout=5, check=True, env=environment
            ).stdout
            language = "tha+eng" if "tha" in langs.splitlines() else "eng"
            completed = subprocess.run(
                ["tesseract", str(source), "stdout", "-l", language],
                capture_output=True, text=True, timeout=30, check=True, env=environment,
            )
            return completed.stdout.strip()
        except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
            return None


def render_pdf_page(document: PdfDocument, index: int, target_scale: float = 2.0) -> Image.Image:
    page = document[index]
    try:
        width, height = page.get_size()
        scale = min(target_scale, sqrt(MAX_IMAGE_PIXELS / max(width * height, 1)))
        bitmap = page.render(scale=scale)
        try:
            return bitmap.to_pil().convert("RGBA")
        finally:
            bitmap.close()
    finally:
        page.close()


def extract_pdf_text(data: bytes) -> tuple[str, list[int]]:
    """Read selectable PDF text and OCR pages that contain only a scan."""
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF มีรหัสผ่าน กรุณาปลดล็อกก่อนเปรียบเทียบ")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f"PDF มีเกิน {MAX_PDF_PAGES} หน้า")
        parts = []
        scanned = []
        for index, page in enumerate(reader.pages):
            value = page.extract_text() or ""
            parts.append(value)
            if not value.strip():
                scanned.append(index)
        if len(scanned) > MAX_PDF_OCR_PAGES:
            raise ValueError(f"PDF มีหน้าสแกนเกิน {MAX_PDF_OCR_PAGES} หน้า กรุณาแบ่งไฟล์ก่อนเปรียบเทียบ")
        if scanned:
            with PdfDocument(data) as document:
                for index in scanned:
                    value = ocr_image(render_pdf_page(document, index))
                    if value is None:
                        raise ValueError("OCR สำหรับ PDF สแกนไม่พร้อมใช้งาน")
                    parts[index] = value
        result = "\n".join(parts)
        if len(result) > MAX_TEXT:
            raise ValueError("ข้อความที่ดึงได้ยาวเกิน 200,000 ตัวอักษร")
        return result, scanned
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("อ่าน PDF ไม่สำเร็จ ตรวจว่าไฟล์ไม่เสียหรือมีรหัสผ่าน") from exc


def source_comparison_text(filename: str, data: bytes, image: Image.Image | None) -> tuple[str | None, bool, list[int]]:
    if image is not None:
        return ocr_image(image), True, []
    if Path(filename).suffix.lower() == ".pdf":
        value, scanned = extract_pdf_text(data)
        return value, bool(scanned), scanned
    return extract_text(filename, data), False, []


def pdf_visual_difference(left_data: bytes, right_data: bytes, scanned_pages: list[int]) -> dict:
    changes = []
    with PdfDocument(left_data) as left, PdfDocument(right_data) as right:
        for index in scanned_pages:
            if index >= len(left) or index >= len(right):
                continue
            visual = visual_difference(render_pdf_page(left, index, 1.5), render_pdf_page(right, index, 1.5))
            if not visual["same"]:
                changes.append({"page": index + 1, **visual})
        return {"same": len(left) == len(right) and not changes, "left_pages": len(left),
                "right_pages": len(right), "checked_pages": [index + 1 for index in scanned_pages],
                "changes": changes}


def visual_difference(left: Image.Image, right: Image.Image) -> dict:
    if left.size != right.size:
        return {"same": False, "left_size": list(left.size), "right_size": list(right.size), "changed_pixels": None, "bounds": None}
    delta = ImageChops.difference(left, right)
    channels = delta.split()
    mask = channels[0]
    for channel in channels[1:]:
        mask = ImageChops.lighter(mask, channel)
    changed = mask.point(lambda value: 255 if value else 0).histogram()[255]
    return {"same": changed == 0, "left_size": list(left.size), "right_size": list(right.size),
            "changed_pixels": changed, "bounds": list(mask.getbbox()) if changed else None}


def position(value: str, offset: int) -> dict:
    return {"line": value.count("\n", 0, offset) + 1, "column": offset - value.rfind("\n", 0, offset)}


def text_difference(left: str, right: str) -> dict:
    if left == right:
        return {"same": True, "left_characters": len(left), "right_characters": len(right), "changes": []}
    # A full character diff can stall on large documents. First align lines, then
    # refine short changed blocks at character level. Equality checks all text.
    limited = False
    if max(len(left), len(right)) <= 3_000:
        operations = [op for op in difflib.SequenceMatcher(None, left, right, autojunk=False).get_opcodes() if op[0] != "equal"]
    else:
        left_lines, right_lines = left.splitlines(keepends=True), right.splitlines(keepends=True)
        if max(len(left_lines), len(right_lines)) <= 2_000:
            left_offsets = [0, *accumulate(map(len, left_lines))]
            right_offsets = [0, *accumulate(map(len, right_lines))]
            line_ops = (op for op in difflib.SequenceMatcher(
                None, left_lines, right_lines, autojunk=False
            ).get_opcodes() if op[0] != "equal")
            operations = []
            for tag, i, j, k, l in line_ops:
                a, b, c, d = left_offsets[i], left_offsets[j], right_offsets[k], right_offsets[l]
                if max(b - a, d - c) <= 3_000:
                    operations.extend((part_tag, a + pi, a + pj, c + pk, c + pl)
                        for part_tag, pi, pj, pk, pl in difflib.SequenceMatcher(
                            None, left[a:b], right[c:d], autojunk=False
                        ).get_opcodes() if part_tag != "equal")
                else:
                    operations.append((tag, a, b, c, d))
                    limited = True
                if len(operations) > 30:
                    limited = True
                    break
        else:
            prefix = 0
            while prefix < min(len(left), len(right)) and left[prefix] == right[prefix]:
                prefix += 1
            suffix = 0
            while (suffix < min(len(left), len(right)) - prefix and left[-1 - suffix] == right[-1 - suffix]):
                suffix += 1
            operations = [("replace", prefix, len(left) - suffix, prefix, len(right) - suffix)]
            limited = True
    changes = []
    for _, i, j, k, l in operations[:30]:
        changes.append({
            "left_at": position(left, i), "right_at": position(right, k),
            "before": left[max(0, i - 60):i],
            "left": left[i:min(j, i + 240)], "right": right[k:min(l, k + 240)],
            "after": left[j:min(len(left), j + 60)],
            "left_truncated": j - i > 240, "right_truncated": l - k > 240,
        })
    return {"same": False, "left_characters": len(left), "right_characters": len(right),
            "changes": changes, "preview_limited": limited or len(operations) > 30}


@app.post("/api/compare")
async def compare(left: UploadFile = File(...), right: UploadFile = File(...)):
    try:
        left_kind, right_kind = kind(left.filename or ""), kind(right.filename or "")
        left_data, right_data = await read_upload(left), await read_upload(right)
        return await run_in_threadpool(
            compare_bytes, left.filename or "", left_data, left_kind,
            right.filename or "", right_data, right_kind,
        )
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


def compare_bytes(left_name: str, left_data: bytes, left_kind: str,
                  right_name: str, right_data: bytes, right_kind: str) -> dict:
    """Run parsing and OCR off the async event loop so pairs can progress together."""
    try:
        left_image = load_image(left_data) if left_kind == "image" else None
        right_image = load_image(right_data) if right_kind == "image" else None
        visual = visual_difference(left_image, right_image) if left_image and right_image else None
        both_spreadsheets = Path(left_name).suffix.lower() in SPREADSHEET_TYPES and Path(right_name).suffix.lower() in SPREADSHEET_TYPES
        spreadsheet = cell_differences(left_name, left_data, right_name, right_data) if both_spreadsheets else None
        left_text, left_ocr, left_scanned = (None, False, []) if both_spreadsheets else source_comparison_text(left_name, left_data, left_image)
        right_text, right_ocr, right_scanned = (None, False, []) if both_spreadsheets else source_comparison_text(right_name, right_data, right_image)
        text = text_difference(left_text, right_text) if left_text is not None and right_text is not None else None
        both_pdf = Path(left_name).suffix.lower() == Path(right_name).suffix.lower() == ".pdf"
        pdf_visual = pdf_visual_difference(left_data, right_data, sorted(set(left_scanned + right_scanned))) if both_pdf and (left_scanned or right_scanned) else None
        if visual and visual["same"]:
            status = "completed"
        elif visual:
            status = "different"
        elif spreadsheet:
            status = "completed" if spreadsheet["same"] else "different"
        elif pdf_visual and not pdf_visual["same"]:
            status = "different"
        elif pdf_visual and pdf_visual["same"] and not left_text and not right_text:
            status = "completed"
        elif left_kind == right_kind == "document" and left_data == right_data:
            status = "completed"
        elif text is None or ((left_image or right_image) and not (left_text or right_text)):
            status = "inconclusive"
        elif not left_text and not right_text:
            status = "inconclusive"
        elif text["same"]:
            status = "matched_ocr" if left_ocr or right_ocr else "completed"
        else:
            status = "different"
        return {"status": status, "left_kind": left_kind, "right_kind": right_kind,
                "visual": visual, "pdf_visual": pdf_visual, "text": text, "spreadsheet": spreadsheet,
                "ocr_available": (left_text is not None and right_text is not None),
                "ocr_language": "tha+eng" if left_ocr or right_ocr else None,
                "ocr_used": left_ocr or right_ocr}
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


def line_report_rows(left: str, right: str):
    left_lines, right_lines = left.splitlines(), right.splitlines()
    if max(len(left_lines), len(right_lines)) > 5_000:
        for number, (old, new) in enumerate(zip_longest(left_lines, right_lines, fillvalue=""), 1):
            if old != new:
                yield ["ข้อความ", "", number, number, old, new]
        return
    matcher = difflib.SequenceMatcher(None, left_lines, right_lines, autojunk=True)
    for tag, i, j, k, l in matcher.get_opcodes():
        if tag == "equal":
            continue
        for offset, (old, new) in enumerate(zip_longest(left_lines[i:j], right_lines[k:l], fillvalue="")):
            yield ["ข้อความ", "", i + offset + 1 if i + offset < j else "",
                   k + offset + 1 if k + offset < l else "", old, new]


def report_bytes(left_name: str, left_data: bytes, right_name: str, right_data: bytes) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["ชนิด", "ชีต / รายละเอียด", "ตำแหน่งต้นฉบับ", "ตำแหน่งฝั่งเทียบ", "ต้นฉบับ", "ฝั่งเทียบ"])
    if Path(left_name).suffix.lower() in SPREADSHEET_TYPES and Path(right_name).suffix.lower() in SPREADSHEET_TYPES:
        left_sheets, left_cells = read_spreadsheet(left_name, left_data)
        right_sheets, right_cells = read_spreadsheet(right_name, right_data)
        for name in left_sheets:
            if name not in right_sheets:
                writer.writerow(["ชีตเฉพาะฝั่งต้นฉบับ", name, "", "", "", ""])
        for name in right_sheets:
            if name not in left_sheets:
                writer.writerow(["ชีตเฉพาะฝั่งเทียบ", name, "", "", "", ""])
        for change in iter_cell_changes(left_cells, right_cells):
            writer.writerow(["เซลล์", change["sheet"], change["cell"], change["cell"], change["left"], change["right"]])
    else:
        left_kind, right_kind = kind(left_name), kind(right_name)
        left_image = load_image(left_data) if left_kind == "image" else None
        right_image = load_image(right_data) if right_kind == "image" else None
        if left_image and right_image:
            visual = visual_difference(left_image, right_image)
            writer.writerow(["รูปภาพ", "ขนาด / พิกเซลต่าง", str(visual["left_size"]), str(visual["right_size"]),
                             visual["changed_pixels"], visual["bounds"]])
        left_text, _, left_scanned = source_comparison_text(left_name, left_data, left_image)
        right_text, _, right_scanned = source_comparison_text(right_name, right_data, right_image)
        if left_text is not None and right_text is not None:
            writer.writerows(line_report_rows(left_text, right_text))
        if Path(left_name).suffix.lower() == Path(right_name).suffix.lower() == ".pdf" and (left_scanned or right_scanned):
            visual = pdf_visual_difference(left_data, right_data, sorted(set(left_scanned + right_scanned)))
            if visual["left_pages"] != visual["right_pages"]:
                writer.writerow(["PDF จำนวนหน้า", "", visual["left_pages"], visual["right_pages"], "", ""])
            for change in visual["changes"]:
                writer.writerow(["PDF ภาพสแกน", f'หน้า {change["page"]}',
                                 change["bounds"], change["bounds"],
                                 change["left_size"], change["right_size"]])
    return output.getvalue().encode("utf-8-sig")


def csv_response(content: bytes) -> Response:
    return Response(content, media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="comparison-report.csv"'})


@app.post("/api/compare/report")
async def compare_report(left: UploadFile = File(...), right: UploadFile = File(...)):
    try:
        kind(left.filename or "")
        kind(right.filename or "")
        left_data, right_data = await read_upload(left), await read_upload(right)
        return csv_response(await run_in_threadpool(report_bytes, left.filename or "", left_data,
                                                    right.filename or "", right_data))
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


class LocalFolders(BaseModel):
    left_path: str
    right_path: str
    match_mode: str = "auto"


class LocalFiles(BaseModel):
    left_path: str
    right_path: str


def require_local_access(request: Request):
    if os.environ.get("LOCAL_FOLDER_ACCESS") != "1" or not request.client or request.client.host not in {"127.0.0.1", "::1"}:
        raise HTTPException(403, detail="การอ่าน path ใช้ได้เฉพาะเซิร์ฟเวอร์ที่รันบนเครื่องนี้และเปิด LOCAL_FOLDER_ACCESS=1")


def local_directory(path: str) -> Path:
    try:
        root = Path(path).expanduser().resolve(strict=True)
        if not root.is_dir():
            raise ValueError("path ที่ระบุไม่ใช่โฟลเดอร์")
        return root
    except OSError as exc:
        raise ValueError("เปิดโฟลเดอร์ไม่ได้ ตรวจ path และสิทธิ์อ่านไฟล์") from exc


def list_local_files(root: Path) -> dict[str, Path]:
    return {item.relative_to(root).as_posix(): item for item in root.rglob("*")
            if item.is_file() and not item.is_symlink() and item.suffix.lower() in SUPPORTED_TYPES}


def natural_path_key(value: str) -> list:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def list_folder_pairs(left_path: str, right_path: str, match_mode: str = "auto") -> dict:
    if match_mode not in {"auto", "name", "order"}:
        raise ValueError("วิธีจับคู่ไม่ถูกต้อง")
    left_root, right_root = local_directory(left_path), local_directory(right_path)
    left_files, right_files = list_local_files(left_root), list_local_files(right_root)
    left_names, right_names = set(left_files), set(right_files)
    exact = sorted(left_names & right_names, key=natural_path_key) if match_mode != "order" else []
    pairs = [(name, name, "name") for name in exact]
    remaining_left = sorted(left_names - set(exact), key=natural_path_key)
    remaining_right = sorted(right_names - set(exact), key=natural_path_key)
    if match_mode != "name":
        pairs.extend((left, right, "order") for left, right in zip(remaining_left, remaining_right))
    pairs.sort(key=lambda pair: natural_path_key(pair[0]))
    return {
        "pairs": [{"name": left, "left_path": str(left_files[left]), "right_path": str(right_files[right]),
                   "left_name": left_files[left].name, "right_name": right_files[right].name,
                   "match": method} for left, right, method in pairs],
        "left_only": len(left_files) - len(pairs),
        "right_only": len(right_files) - len(pairs),
        "exact_pairs": sum(method == "name" for _, _, method in pairs),
        "ordered_pairs": sum(method == "order" for _, _, method in pairs),
    }


def read_local_file(path: str) -> tuple[str, bytes]:
    try:
        source = Path(path).expanduser()
        if source.is_symlink() or not source.is_file():
            raise ValueError("path ที่ระบุไม่ใช่ไฟล์ปกติ")
        kind(source.name)
        if source.stat().st_size > MAX_BYTES:
            raise ValueError("ไฟล์ใหญ่เกิน 20 MB")
        return source.name, source.read_bytes()
    except OSError as exc:
        raise ValueError("อ่านไฟล์จากโฟลเดอร์ไม่ได้") from exc


@app.post("/api/local/list")
async def local_list(request: Request, folders: LocalFolders):
    require_local_access(request)
    try:
        return await run_in_threadpool(list_folder_pairs, folders.left_path, folders.right_path, folders.match_mode)
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


@app.post("/api/local/compare")
async def local_compare(request: Request, files: LocalFiles):
    require_local_access(request)
    try:
        left_name, left_data = await run_in_threadpool(read_local_file, files.left_path)
        right_name, right_data = await run_in_threadpool(read_local_file, files.right_path)
        return await run_in_threadpool(compare_bytes, left_name, left_data, kind(left_name),
                                       right_name, right_data, kind(right_name))
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


@app.post("/api/local/report")
async def local_report(request: Request, files: LocalFiles):
    require_local_access(request)
    try:
        left_name, left_data = await run_in_threadpool(read_local_file, files.left_path)
        right_name, right_data = await run_in_threadpool(read_local_file, files.right_path)
        return csv_response(await run_in_threadpool(report_bytes, left_name, left_data, right_name, right_data))
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


def read_json_object(raw: str, label: str) -> dict:
    try:
        value = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{label} ไม่ใช่ JSON ที่ถูกต้อง") from exc
    if not isinstance(value, dict):
        raise ValueError(f"{label} ต้องเป็น object")
    return value


def source_text(filename: str, data: bytes) -> str:
    if kind(filename) == "image":
        value = ocr_image(load_image(data))
        if value is None:
            raise ValueError("OCR ไม่พร้อมใช้งานสำหรับรูปภาพนี้")
        return value
    return extract_text(filename, data)


def extract_source(filename: str, data: bytes, keywords: list[str]) -> dict:
    hits = extract_hits(filename, data, keywords, source_text)
    return {"source": filename, "hits": hits, "hit_count": len(hits)}


@app.post("/api/keywords/extract")
async def keywords_extract(source: UploadFile = File(...), keywords_json: str = Form(...)):
    try:
        kind(source.filename or "")
        keywords = parse_keywords(json.loads(keywords_json))
        data = await read_upload(source)
        return await run_in_threadpool(extract_source, source.filename or "", data, keywords)
    except (ValueError, TypeError, json.JSONDecodeError) as exc:
        raise HTTPException(422, detail=str(exc)) from exc


class LocalPath(BaseModel):
    path: str


class LocalKeywords(BaseModel):
    path: str
    keywords: list[str]


def mac_choose_directory() -> str:
    if sys.platform != "darwin":
        raise ValueError("ปุ่มเลือกโฟลเดอร์ใช้ได้เมื่อรันเซิร์ฟเวอร์บน Mac เท่านั้น")
    try:
        chosen = subprocess.run(
            ["osascript", "-e", 'POSIX path of (choose folder with prompt "เลือกโฟลเดอร์ที่จะอ่านไฟล์")'],
            capture_output=True, text=True, timeout=120, check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise ValueError("เปิดหน้าต่างเลือกโฟลเดอร์ไม่สำเร็จ") from exc
    if chosen.returncode or not chosen.stdout.strip():
        raise ValueError("ไม่ได้เลือกโฟลเดอร์")
    return str(local_directory(chosen.stdout.strip()))


@app.post("/api/local/choose-directory")
async def local_choose_directory(request: Request):
    require_local_access(request)
    try:
        return {"path": await run_in_threadpool(mac_choose_directory)}
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


@app.post("/api/local/source-list")
async def local_source_list(request: Request, payload: LocalPath):
    require_local_access(request)
    try:
        files = await run_in_threadpool(list_local_files, local_directory(payload.path))
        return {"sources": [{"name": name, "path": str(path)} for name, path in sorted(files.items())]}
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


@app.post("/api/local/keywords/extract")
async def local_keywords_extract(request: Request, payload: LocalKeywords):
    require_local_access(request)
    try:
        keywords = parse_keywords(payload.keywords)
        filename, data = await run_in_threadpool(read_local_file, payload.path)
        return await run_in_threadpool(extract_source, filename, data, keywords)
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


class WebGeneration(BaseModel):
    template: dict
    values: dict
    format: str


def document_response(content: bytes, mime: str, extension: str) -> Response:
    return Response(content, media_type=mime,
                    headers={"Content-Disposition": f'attachment; filename="generated-document.{extension}"'})


@app.post("/api/generate/web")
async def generate_web(payload: WebGeneration):
    try:
        content, mime = await run_in_threadpool(render_web_template, payload.template, payload.values, payload.format)
        return document_response(content, mime, payload.format)
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc


@app.post("/api/generate/file")
async def generate_file(template: UploadFile = File(...), values_json: str = Form(...)):
    try:
        values = read_json_object(values_json, "ข้อมูล keyword")
        data = await read_upload(template)
        content, mime, extension = await run_in_threadpool(render_file_template, template.filename or "", data, values)
        return document_response(content, mime, extension)
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc
