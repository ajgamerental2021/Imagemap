import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter
from xlwt import Workbook as LegacyWorkbook

from app import app, extract_text


client = TestClient(app)


def request(left_name, left_bytes, right_name, right_bytes):
    return client.post('/api/compare', files={
        'left': (left_name, left_bytes), 'right': (right_name, right_bytes),
    })


def image_bytes(text=None, color='white'):
    image = Image.new('RGB', (430, 105), color)
    if text:
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 42)
        draw.text((12, 22), text, fill='black', font=font)
    output = io.BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


class ComparisonTests(unittest.TestCase):
    def test_exact_text_and_multiple_changes(self):
        same = request('a.txt', b'alpha\nbeta', 'b.txt', b'alpha\nbeta')
        self.assertEqual(same.status_code, 200)
        self.assertEqual(same.json()['status'], 'completed')
        different = request('a.txt', b'alpha\nbeta', 'b.txt', b'alpha\ngamma')
        body = different.json()
        self.assertEqual(body['status'], 'different')
        self.assertEqual(body['text']['changes'][0]['left_at']['line'], 2)

    def test_large_multiline_document_shows_separate_changes(self):
        lines = [f'line {i:04d}: value\n' for i in range(400)]
        changed = lines.copy()
        changed[12] = 'line 0012: other\n'
        changed[310] = 'line 0310: other\n'
        body = request('a.txt', ''.join(lines).encode(), 'b.txt', ''.join(changed).encode()).json()
        self.assertEqual(body['status'], 'different')
        positions = [item['left_at']['line'] for item in body['text']['changes']]
        self.assertIn(13, positions)
        self.assertIn(311, positions)

    def test_identical_and_different_pixels(self):
        white = image_bytes()
        changed = image_bytes(color='lightgray')
        self.assertEqual(request('a.png', white, 'b.png', white).json()['status'], 'completed')
        body = request('a.png', white, 'b.png', changed).json()
        self.assertEqual(body['status'], 'different')
        self.assertGreater(body['visual']['changed_pixels'], 0)

    def test_avif_image_upload(self):
        image = Image.new('RGB', (20, 20), 'white')
        output = io.BytesIO()
        image.save(output, format='AVIF')
        self.assertEqual(request('image.png.avif', output.getvalue(), 'copy.avif', output.getvalue()).json()['status'], 'completed')

    def test_image_to_text_uses_ocr_without_claiming_exact_certainty(self):
        body = request('image.png', image_bytes('HELLO 123'), 'text.txt', b'HELLO 123').json()
        self.assertEqual(body['status'], 'matched_ocr')
        self.assertTrue(body['text']['same'])

    def test_unreadable_pdf_text_is_inconclusive(self):
        def pdf(width):
            writer = PdfWriter()
            writer.add_blank_page(width=width, height=200)
            buffer = io.BytesIO()
            writer.write(buffer)
            return buffer.getvalue()
        self.assertEqual(request('a.pdf', pdf(100), 'b.pdf', pdf(110)).json()['status'], 'inconclusive')

    def test_docx_and_xlsx_extraction(self):
        doc = Document()
        doc.add_paragraph('Invoice 42')
        output = io.BytesIO()
        doc.save(output)
        self.assertEqual(request('a.docx', output.getvalue(), 'b.txt', b'Invoice 42').json()['status'], 'completed')
        book = Workbook()
        book.active.title = 'Sheet1'
        book.active['A1'] = 'Total'
        output = io.BytesIO()
        book.save(output)
        self.assertEqual(request('a.xlsx', output.getvalue(), 'b.txt', '[ชีต: Sheet1]\nTotal'.encode()).json()['status'], 'completed')

    def test_excel_xls_xlsx_and_xlsm_compare_cell_values(self):
        legacy = LegacyWorkbook()
        sheet = legacy.add_sheet('Sheet1')
        sheet.write(0, 0, 'Total')
        sheet.write(0, 1, 42)
        xls = io.BytesIO()
        legacy.save(xls)
        modern = Workbook()
        modern.active.title = 'Sheet1'
        modern.active.append(['Total', 42])
        xlsx = io.BytesIO()
        modern.save(xlsx)
        self.assertEqual(request('old.xls', xls.getvalue(), 'new.xlsx', xlsx.getvalue()).json()['status'], 'completed')
        self.assertEqual(request('book.xlsm', xlsx.getvalue(), 'old.xls', xls.getvalue()).json()['status'], 'completed')

    def test_excel_reports_exact_cells_and_full_csv(self):
        def workbook(values):
            book = Workbook()
            book.active.title = 'Orders'
            for row in values:
                book.active.append(row)
            output = io.BytesIO()
            book.save(output)
            return output.getvalue()
        left = workbook([['Item', 'Amount'], ['A', 10], ['B', 20]])
        right = workbook([['Item', 'Amount'], ['A', 11], ['B', None], ['C', 30]])
        result = request('left.xlsx', left, 'right.xlsx', right).json()
        self.assertEqual(result['status'], 'different')
        self.assertEqual(result['spreadsheet']['total_changes'], 4)
        self.assertEqual([change['cell'] for change in result['spreadsheet']['changes']], ['B2', 'B3', 'A4', 'B4'])
        report = client.post('/api/compare/report', files={
            'left': ('left.xlsx', left), 'right': ('right.xlsx', right),
        })
        self.assertEqual(report.status_code, 200)
        self.assertIn('B2', report.content.decode('utf-8-sig'))
        self.assertIn('B4', report.content.decode('utf-8-sig'))

    def test_local_folder_mode_reads_files_without_upload(self):
        local_client = TestClient(app, client=('127.0.0.1', 50000))
        with tempfile.TemporaryDirectory() as directory:
            left = Path(directory) / 'original'
            right = Path(directory) / 'revised'
            left.mkdir(); right.mkdir()
            (left / 'item.txt').write_text('old', encoding='utf-8')
            (right / 'item.txt').write_text('new', encoding='utf-8')
            payload = {'left_path': str(left), 'right_path': str(right)}
            self.assertEqual(local_client.post('/api/local/list', json=payload).status_code, 403)
            with patch.dict(os.environ, {'LOCAL_FOLDER_ACCESS': '1'}):
                listing = local_client.post('/api/local/list', json=payload)
                self.assertEqual(listing.status_code, 200)
                self.assertEqual(len(listing.json()['pairs']), 1)
                pair = listing.json()['pairs'][0]
                comparison = local_client.post('/api/local/compare', json=pair)
                self.assertEqual(comparison.status_code, 200)
                self.assertEqual(comparison.json()['status'], 'different')
                report = local_client.post('/api/local/report', json=pair)
                self.assertIn('old', report.content.decode('utf-8-sig'))

    def test_xlsb_extracts_sheet_and_cell_values(self):
        sample = (Path(__file__).parent / 'fixtures' / 'sample.xlsb').read_bytes()
        content = extract_text('sample.xlsb', sample)
        self.assertTrue(content.startswith('[ชีต: Test]\nA\tB\tA\tB'))
        self.assertIn('42.1337', content)
        self.assertEqual(request('sample.xlsb', sample, 'copy.xlsb', sample).json()['status'], 'completed')

    def test_rejects_unsupported_and_invalid_input(self):
        unsupported = request('a.exe', b'123', 'b.txt', b'123')
        self.assertEqual(unsupported.status_code, 422)
        self.assertEqual(request('a.txt', b'\xff', 'b.txt', b'abc').status_code, 422)
        self.assertEqual(request('a.txt', b'', 'b.txt', b'abc').status_code, 422)


if __name__ == '__main__':
    unittest.main()
