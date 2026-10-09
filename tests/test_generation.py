"""End-to-end keyword extraction and document rendering checks."""

import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app import app

client = TestClient(app)


def workbook_bytes(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class GenerationTests(unittest.TestCase):
    def test_extract_multiple_sources_and_generate_each_format(self):
        keywords = '["ชื่อลูกค้า", "ยอดรวม"]'
        first = client.post('/api/keywords/extract', files={'source': ('a.txt', 'ชื่อลูกค้า: สมหญิง\nยอดรวม: 100'.encode())},
                            data={'keywords_json': keywords})
        second = client.post('/api/keywords/extract', files={'source': ('b.xlsx', workbook_bytes([
            ['ชื่อลูกค้า', 'สมชาย'], ['ยอดรวม', 200], ['ชื่อลูกค้า', 'มานี']]))}, data={'keywords_json': keywords})
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertIn('Sheet!A1', [hit['location'] for hit in second.json()['hits']])
        hits = first.json()['hits'] + second.json()['hits']
        values = {keyword: [hit['value'] for hit in hits if hit['keyword'] == keyword] for keyword in ['ชื่อลูกค้า', 'ยอดรวม']}
        spec = {'title': 'รายงาน {{ชื่อลูกค้า:count}} ราย', 'blocks': [
            {'type': 'field', 'label': 'รายชื่อ', 'text': '{{ชื่อลูกค้า}}'},
            {'type': 'paragraph', 'text': 'รวม {{ยอดรวม:first}} / {{ยอดรวม:last}}'}]}
        for fmt in ('docx', 'xlsx', 'pdf'):
            with self.subTest(fmt=fmt):
                response = client.post('/api/generate/web', json={'template': spec, 'values': values, 'format': fmt})
                self.assertEqual(response.status_code, 200, response.text[:300])
                self.assertGreater(len(response.content), 1000)
                if fmt == 'docx':
                    doc = Document(io.BytesIO(response.content))
                    body = '\n'.join(p.text for p in doc.paragraphs)
                    self.assertIn('สมหญิง\nสมชาย\nมานี', body)
                elif fmt == 'xlsx':
                    book = load_workbook(io.BytesIO(response.content))
                    self.assertEqual([book.active[f'B{row}'].value for row in (2, 3, 4)], ['สมหญิง', 'สมชาย', 'มานี'])
                else:
                    self.assertEqual(len(PdfReader(io.BytesIO(response.content)).pages), 1)

        long_values = {'name': [f'Person {number}' for number in range(150)]}
        long_spec = {'title': 'People', 'blocks': [{'type': 'field', 'label': 'Names', 'text': '{{name}}'}]}
        pdf = client.post('/api/generate/web', json={'template': long_spec, 'values': long_values, 'format': 'pdf'})
        self.assertEqual(pdf.status_code, 200)
        self.assertGreater(len(PdfReader(io.BytesIO(pdf.content)).pages), 1)
        excel = client.post('/api/generate/web', json={'template': long_spec, 'values': long_values, 'format': 'xlsx'})
        self.assertEqual(load_workbook(io.BytesIO(excel.content)).active['B151'].value, 'Person 149')

    def test_uploaded_word_excel_and_pdf_templates(self):
        values = {'name': ['Alice', 'Bob']}
        word = Document()
        paragraph = word.add_paragraph()
        paragraph.add_run('Dear {{na').bold = True
        paragraph.add_run('me:first}}')
        word_bytes = io.BytesIO(); word.save(word_bytes)
        response = client.post('/api/generate/file', files={'template': ('template.docx', word_bytes.getvalue())},
                               data={'values_json': '{"name":["Alice","Bob"]}'})
        self.assertEqual(response.status_code, 200, response.text[:300])
        self.assertEqual(Document(io.BytesIO(response.content)).paragraphs[0].text, 'Dear Alice')

        excel = workbook_bytes([['Clients', '{{name}}'], ['Count', '{{name:count}}']])
        response = client.post('/api/generate/file', files={'template': ('template.xlsx', excel)},
                               data={'values_json': '{"name":["Alice","Bob"]}'})
        self.assertEqual(response.status_code, 200)
        filled = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(filled['B1'].value, 'Alice\nBob')
        self.assertEqual(filled['B2'].value, '2')

        pdf = io.BytesIO()
        page = canvas.Canvas(pdf)
        page.acroForm.textfield(name='name', x=80, y=700, width=200, height=20)
        page.showPage()
        page.save()
        response = client.post('/api/generate/file', files={'template': ('template.pdf', pdf.getvalue())},
                               data={'values_json': '{"name":["Alice","Bob"]}'})
        self.assertEqual(response.status_code, 200, response.text[:300])
        self.assertEqual(PdfReader(io.BytesIO(response.content)).get_fields()['name']['/V'], 'Alice\nBob')

    def test_local_source_directory_and_bad_inputs(self):
        local_client = TestClient(app, client=('127.0.0.1', 54321))
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {'LOCAL_FOLDER_ACCESS': '1'}):
            source = Path(directory) / 'a.txt'
            source.write_text('Order: 123', encoding='utf-8')
            listing = local_client.post('/api/local/source-list', json={'path': directory})
            self.assertEqual(len(listing.json()['sources']), 1)
            extracted = local_client.post('/api/local/keywords/extract', json={'path': str(source), 'keywords': ['Order']})
            self.assertEqual(extracted.json()['hits'][0]['value'], '123')
            with patch('app.sys.platform', 'darwin'), patch('app.subprocess.run') as picker:
                picker.return_value.returncode = 0
                picker.return_value.stdout = directory + '\n'
                chosen = local_client.post('/api/local/choose-directory')
                self.assertEqual(chosen.json()['path'], directory)
        invalid = client.post('/api/keywords/extract', files={'source': ('a.txt', b'Order: 1')},
                              data={'keywords_json': '{"bad":true}'})
        self.assertEqual(invalid.status_code, 422)
        plain = io.BytesIO()
        page = canvas.Canvas(plain)
        page.drawString(40, 700, 'No form fields')
        page.save()
        missing = client.post('/api/generate/file', files={'template': ('template.pdf', plain.getvalue())},
                              data={'values_json': '{"Order":["123"]}'})
        self.assertEqual(missing.status_code, 422)


if __name__ == '__main__':
    unittest.main()
