"""Structured spreadsheet data shared by comparisons and document generation."""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook
from pyxlsb import open_workbook as open_xlsb
from xlrd import open_workbook as open_xls

SPREADSHEET_TYPES = {".xlsx", ".xlsm", ".xls", ".xlsb"}


def cell_text(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_spreadsheet(filename: str, data: bytes) -> tuple[list[str], dict[tuple[str, int, int], str]]:
    """Return sheet names and nonempty cell values, with one-based coordinates."""
    suffix = Path(filename).suffix.lower()
    sheets: list[str] = []
    cells: dict[tuple[str, int, int], str] = {}
    try:
        if suffix in {".xlsx", ".xlsm"}:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                for sheet in book:
                    sheets.append(sheet.title)
                    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        for column_number, value in enumerate(row, 1):
                            rendered = cell_text(value)
                            if rendered:
                                cells[(sheet.title, row_number, column_number)] = rendered
            finally:
                book.close()
        elif suffix == ".xls":
            book = open_xls(file_contents=data)
            for sheet in book.sheets():
                sheets.append(sheet.name)
                for row_number in range(sheet.nrows):
                    for column_number, value in enumerate(sheet.row_values(row_number), 1):
                        rendered = cell_text(value)
                        if rendered:
                            cells[(sheet.name, row_number + 1, column_number)] = rendered
        elif suffix == ".xlsb":
            with open_xlsb(io.BytesIO(data)) as book:
                for name in book.sheets:
                    sheets.append(name)
                    with book.get_sheet(name) as sheet:
                        for row in sheet.rows(sparse=True):
                            for cell in row:
                                rendered = cell_text(cell.v)
                                if rendered:
                                    cells[(name, cell.r + 1, cell.c + 1)] = rendered
        else:
            raise ValueError("ไม่ใช่ไฟล์ Excel ที่รองรับ")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("อ่านข้อมูลเซลล์ Excel ไม่สำเร็จ") from exc
    return sheets, cells


def column_label(number: int) -> str:
    output = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        output = chr(65 + remainder) + output
    return output


def iter_cell_changes(left_cells: dict, right_cells: dict):
    for sheet, row, column in sorted(left_cells.keys() | right_cells.keys()):
        left = left_cells.get((sheet, row, column), "")
        right = right_cells.get((sheet, row, column), "")
        if left == right:
            continue
        yield {
            "sheet": sheet,
            "cell": f"{column_label(column)}{row}",
            "left": left,
            "right": right,
            "type": "changed" if left and right else "removed" if left else "added",
        }


def cell_differences(
    left_name: str, left_data: bytes, right_name: str, right_data: bytes,
    preview_limit: int = 500,
) -> dict:
    left_sheets, left_cells = read_spreadsheet(left_name, left_data)
    right_sheets, right_cells = read_spreadsheet(right_name, right_data)
    sheet_changes = [
        {"sheet": sheet, "side": "left"} for sheet in left_sheets if sheet not in right_sheets
    ] + [
        {"sheet": sheet, "side": "right"} for sheet in right_sheets if sheet not in left_sheets
    ]
    changes = []
    total = 0
    for change in iter_cell_changes(left_cells, right_cells):
        total += 1
        if len(changes) < preview_limit:
            changes.append(change)
    return {
        "same": total == 0 and not sheet_changes,
        "total_changes": total,
        "changes": changes,
        "sheet_changes": sheet_changes,
        "preview_limited": total > len(changes),
        "left_cells": len(left_cells),
        "right_cells": len(right_cells),
    }
